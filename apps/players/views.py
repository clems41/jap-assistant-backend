import datetime
import functools
import io
import re
from dataclasses import dataclass

import openpyxl
from django.db import transaction
from django.shortcuts import get_object_or_404
from drf_spectacular.utils import OpenApiParameter, extend_schema
from rest_framework import generics, status
from rest_framework.exceptions import ValidationError
from rest_framework.parsers import MultiPartParser
from rest_framework.permissions import AllowAny, IsAuthenticated
from rest_framework.request import Request
from rest_framework.response import Response
from rest_framework.views import APIView

from apps.common.exceptions import ConflictError
from apps.notifications.services import Resource, notify_public_update
from apps.tournaments.models import Tournament
from apps.tournaments.views import PublicTournamentScopedMixin

from .models import Pair, Player
from .serializers import PairImportSerializer, PairSerializer, PublicPairSerializer
from .services.ranking_matching_service import match_and_update_rankings

SHEET_NAME = "Tableau final"

EXPECTED_HEADERS = [
    "Nom de l'épreuve",
    "Catégorie de l'épreuve",
    "Position de la paire",
    "Nom joueur 1",
    "Prénom joueur 1",
    "Date de naissance joueur 1",
    "Licence joueur 1",
    "Club joueur 1",
    "Classement joueur 1",
    "Mail joueur 1",
    "Téléphone joueur 1",
    "Nom joueur 2",
    "Prénom joueur 2",
    "Date de naissance joueur 2",
    "Licence joueur 2",
    "Club joueur 2",
    "Classement joueur 2",
    "Mail joueur 2",
    "Téléphone joueur 2",
    "Poids paire",
]

_LICENSE_SEASON_SUFFIX_RE = re.compile(r"\s*\(\d{4}\)\s*$")


@dataclass
class ParsedPairRow:
    row_number: int
    player1_license: str
    player1_defaults: dict
    player2_license: str
    player2_defaults: dict
    weight: float | None


def _value_to_str(value: object) -> str:
    """Normalize a raw openpyxl cell value to a stripped string.

    A numeric value that holds an integer-looking float (e.g. a license
    number or ranking exported as an Excel number) must not produce a
    trailing ".0" artifact.
    """
    if value is None:
        return ""
    if isinstance(value, float):
        if value == int(value):
            return str(int(value))
        return str(value)
    return str(value).strip()


def _value_to_int(value: object) -> int | None:
    raw = _value_to_str(value)
    if not raw:
        return None
    return int(float(raw))


def _value_to_float(value: object) -> float | None:
    raw = _value_to_str(value)
    if not raw:
        return None
    return float(raw)


def _value_to_date(value: object) -> tuple[datetime.date | None, str | None]:
    """Return (date, error). Supports native datetime values and a JJ/MM/AAAA
    text fallback."""
    if value is None or value == "":
        return None, None
    if isinstance(value, datetime.datetime):
        return value.date(), None
    if isinstance(value, datetime.date):
        return value, None
    raw = _value_to_str(value)
    if not raw:
        return None, None
    try:
        day, month, year = raw.split("/")
        return datetime.date(int(year), int(month), int(day)), None
    except (ValueError, TypeError):
        return None, "date de naissance invalide"


def _strip_license_season_suffix(raw: str) -> str:
    """Strip a trailing ' (AAAA)' season annotation from a license code.

    e.g. "3271896H (2026)" -> "3271896H". A license number is the upsert key
    for Player, so this must be stripped to keep the same physical player
    identified across seasons on re-import.
    """
    return _LICENSE_SEASON_SUFFIX_RE.sub("", raw).strip()


def _build_header_index(header_row: tuple) -> dict[str, int]:
    """Resolve each expected header to its actual column index (lookup by name)."""
    actual = [_value_to_str(value) for value in header_row]
    index: dict[str, int] = {}
    for header in EXPECTED_HEADERS:
        if header not in actual:
            raise ValueError(header)
        index[header] = actual.index(header)
    return index


def _build_player_defaults(
    row: tuple, col: dict[str, int], suffix: str
) -> tuple[str, dict, str | None]:
    """Build (license_number, defaults, error) for one player (joueur 1 or
    joueur 2) of a row.

    `suffix` is " joueur 1" or " joueur 2". Mirrors the field-by-field
    priority rules:
    - last_name/first_name/club/email/phone: always overwritten from file (even blank).
    - ranking/birth_date: omitted from defaults when the cell is empty, to
      preserve whatever the DB already has.
    """
    last_name = _value_to_str(row[col[f"Nom{suffix}"]])
    first_name = _value_to_str(row[col[f"Prénom{suffix}"]])
    club = _value_to_str(row[col[f"Club{suffix}"]])
    email = _value_to_str(row[col[f"Mail{suffix}"]])
    phone = _value_to_str(row[col[f"Téléphone{suffix}"]])

    license_raw = _value_to_str(row[col[f"Licence{suffix}"]])
    license_number = _strip_license_season_suffix(license_raw)

    defaults: dict = {
        "last_name": last_name,
        "first_name": first_name,
        "club": club,
        "email": email,
        "phone": phone,
    }

    ranking = _value_to_int(row[col[f"Classement{suffix}"]])
    if ranking is not None:
        defaults["ranking"] = ranking

    birth_date, date_error = _value_to_date(row[col[f"Date de naissance{suffix}"]])
    if date_error:
        return license_number, defaults, date_error
    if birth_date is not None:
        defaults["birth_date"] = birth_date

    return license_number, defaults, None


def _parse_xlsx(content: bytes) -> tuple[list[ParsedPairRow], str | None]:
    """Parse the uploaded .xlsx bytes into a list of ParsedPairRow.

    Returns (rows, error). On error, rows is [] and error is a French message
    suitable for ValidationError.
    """
    try:
        workbook = openpyxl.load_workbook(
            io.BytesIO(content), data_only=True, read_only=True
        )
    except Exception:
        return [], "Le fichier est illisible ou corrompu."

    if SHEET_NAME not in workbook.sheetnames:
        return [], f"La feuille '{SHEET_NAME}' est introuvable dans le fichier."

    sheet = workbook[SHEET_NAME]
    data_rows = sheet.iter_rows(values_only=True)

    header_row = next(data_rows, None)
    if header_row is None:
        return [], "Le fichier ne contient aucune ligne d'en-tête."

    try:
        col = _build_header_index(header_row)
    except ValueError as exc:
        return (
            [],
            f"En-têtes invalides. Colonne manquante : '{exc.args[0]}'.",
        )

    rows: list[ParsedPairRow] = []
    seen_licenses: set[str] = set()

    for row_number, row in enumerate(data_rows, start=2):
        if all(value is None or value == "" for value in row):
            continue

        p1_license, p1_defaults, p1_error = _build_player_defaults(
            row, col, suffix=" joueur 1"
        )
        if p1_error:
            return [], f"Ligne {row_number} : {p1_error} (joueur 1)."

        p2_license, p2_defaults, p2_error = _build_player_defaults(
            row, col, suffix=" joueur 2"
        )
        if p2_error:
            return [], f"Ligne {row_number} : {p2_error} (joueur 2)."

        for lic, player_label in [(p1_license, "joueur 1"), (p2_license, "joueur 2")]:
            if not lic:
                return (
                    [],
                    f"Ligne {row_number} : le numéro de licence est obligatoire ({player_label}).",
                )
            if lic in seen_licenses:
                return (
                    [],
                    f"Ligne {row_number} : le numéro de licence '{lic}' est en doublon dans le fichier.",
                )
            seen_licenses.add(lic)

        weight = _value_to_float(row[col["Poids paire"]])

        rows.append(
            ParsedPairRow(
                row_number=row_number,
                player1_license=p1_license,
                player1_defaults=p1_defaults,
                player2_license=p2_license,
                player2_defaults=p2_defaults,
                weight=weight,
            )
        )

    return rows, None


def _create_pairs_from_xlsx(
    tournament: Tournament, rows: list[ParsedPairRow]
) -> list[Pair]:
    with transaction.atomic():
        pairs = []
        for row in rows:
            player1, _ = Player.objects.update_or_create(
                license_number=row.player1_license,
                defaults=row.player1_defaults,
            )
            player2, _ = Player.objects.update_or_create(
                license_number=row.player2_license,
                defaults=row.player2_defaults,
            )
            pair, created = Pair.objects.get_or_create(
                tournament=tournament,
                player1=player1,
                player2=player2,
                defaults={"weight": row.weight},
            )
            if not created and row.weight is not None:
                pair.weight = row.weight
                pair.save(update_fields=["weight", "updated_at"])
            pairs.append(pair)
    return pairs


class TournamentScopedMixin:
    request: Request
    kwargs: dict

    @functools.cached_property
    def _tournament(self) -> Tournament:
        return get_object_or_404(
            Tournament, pk=self.kwargs["tournament_id"], owner=self.request.user
        )


class PairListCreateView(TournamentScopedMixin, generics.ListCreateAPIView):
    serializer_class = PairSerializer
    permission_classes = [IsAuthenticated]
    pagination_class = None
    filter_backends = []

    def get_queryset(self):
        return Pair.objects.filter(tournament=self._tournament).select_related(
            "player1", "player2"
        )

    def get_serializer_context(self):
        ctx = super().get_serializer_context()
        ctx["tournament"] = self._tournament
        return ctx

    def perform_create(self, serializer):
        if self._tournament.is_locked:
            raise ConflictError("Les paires ne peuvent plus être modifiées.")
        serializer.save(tournament=self._tournament)
        notify_public_update(self._tournament, Resource.PAIRS, Resource.TOURNAMENT)


class PublicPairListView(PublicTournamentScopedMixin, generics.ListAPIView):
    """List all pairs of a tournament — public, unauthenticated, PII-free."""

    serializer_class = PublicPairSerializer
    permission_classes = [AllowAny]
    pagination_class = None
    filter_backends = []

    def get_queryset(self):
        return Pair.objects.filter(tournament=self._tournament).select_related(
            "player1", "player2"
        )


class PairDetailView(TournamentScopedMixin, generics.RetrieveUpdateDestroyAPIView):
    serializer_class = PairSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        return Pair.objects.filter(tournament=self._tournament).select_related(
            "player1", "player2"
        )

    def get_serializer_context(self):
        ctx = super().get_serializer_context()
        ctx["tournament"] = self._tournament
        return ctx

    def perform_update(self, serializer):
        if self._tournament.is_locked:
            raise ConflictError("Les paires ne peuvent plus être modifiées.")
        serializer.save()
        notify_public_update(self._tournament, Resource.PAIRS, Resource.TOURNAMENT)

    def perform_destroy(self, instance):
        if self._tournament.is_locked:
            raise ConflictError("Les paires ne peuvent plus être modifiées.")
        instance.delete()
        notify_public_update(self._tournament, Resource.PAIRS, Resource.TOURNAMENT)


class PairImportView(TournamentScopedMixin, APIView):
    parser_classes = [MultiPartParser]
    permission_classes = [IsAuthenticated]

    @extend_schema(
        request=PairImportSerializer,
        responses={201: PairSerializer(many=True)},
        parameters=[
            OpenApiParameter(
                name="tournament_id",
                location=OpenApiParameter.PATH,
                type=int,
            ),
        ],
        summary="Import des paires depuis un fichier Excel (.xlsx)",
        description=(
            "Importe les paires d'un tournoi depuis un fichier Excel (.xlsx), "
            "feuille 'Tableau final' (une ligne = une paire). "
            "Les paires existantes non présentes dans le fichier sont conservées. "
            "Si une paire avec les mêmes joueurs existe déjà, son poids est mis à jour. "
            "Les joueurs identifiés par leur numéro de licence sont créés ou mis à jour. "
            "Après l'import, un matching FFT est automatiquement déclenché pour les joueurs "
            "sans classement. "
            "La réponse contient l'ensemble des paires du tournoi (pas uniquement celles du "
            "fichier), avec les classements et poids mis à jour."
        ),
    )
    def post(self, request: Request, tournament_id: int) -> Response:
        tournament = self._tournament

        if tournament.is_locked:
            raise ConflictError("Les paires ne peuvent plus être modifiées.")

        file_serializer = PairImportSerializer(data=request.data)
        file_serializer.is_valid(raise_exception=True)

        uploaded_file = file_serializer.validated_data["file"]
        content = uploaded_file.read()
        rows, error = _parse_xlsx(content)
        if error:
            raise ValidationError({"file": [error]})

        _create_pairs_from_xlsx(tournament, rows)
        pairs = match_and_update_rankings(tournament)
        output = PairSerializer(
            pairs,
            many=True,
            context={"request": request, "tournament": tournament},
        )

        notify_public_update(tournament, Resource.PAIRS, Resource.TOURNAMENT)

        return Response(output.data, status=status.HTTP_201_CREATED)


class RankingMatchingView(TournamentScopedMixin, APIView):
    permission_classes = [IsAuthenticated]

    @extend_schema(
        responses={200: PairSerializer(many=True)},
        summary="Matching des classements FFT pour les paires du tournoi",
        description=(
            "Pour chaque joueur sans classement, cherche une correspondance dans la table FFTRanking "
            "par nom/prénom (filtré par genre du tournoi). "
            "Les classements existants ne sont pas modifiés. "
            "Recalcule le poids de la paire si les deux joueurs ont un classement."
        ),
        parameters=[
            OpenApiParameter(
                name="tournament_id",
                location=OpenApiParameter.PATH,
                type=int,
            ),
        ],
    )
    def get(self, request: Request, tournament_id: int) -> Response:
        if self._tournament.is_locked:
            raise ConflictError("Les paires ne peuvent plus être modifiées.")
        pairs = match_and_update_rankings(self._tournament)
        serializer = PairSerializer(
            pairs,
            many=True,
            context={"request": request, "tournament": self._tournament},
        )

        notify_public_update(self._tournament, Resource.PAIRS, Resource.TOURNAMENT)

        return Response(serializer.data)
